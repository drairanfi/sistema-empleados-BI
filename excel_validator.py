"""Validates the structure and content of the employees Excel file."""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from config import ALLOWED_EXTENSIONS, EXPECTED_COLUMNS, VALID_DOC_TYPES


class ExcelValidationError(Exception):
    """Raised when the Excel file does not meet the expected format."""

    def __init__(self, errors):
        self.errors = errors
        super().__init__("\n".join(errors))


def _is_blank(value):
    return value is None or str(value).strip() == ""


def _is_number(value):
    if isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    try:
        float(str(value).strip())
        return True
    except ValueError:
        return False


def normalize_doc_number(value):
    # Excel stores numeric documents as floats (1076669247.0)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def validate_file_path(path):
    file = Path(path)
    if not file.exists():
        raise ExcelValidationError([f"No se encontró el archivo: {path}"])
    if file.suffix.lower() not in ALLOWED_EXTENSIONS:
        raise ExcelValidationError(
            [f"Extensión '{file.suffix}' no válida. Permitidas: {', '.join(ALLOWED_EXTENSIONS)}"]
        )
    if file.stat().st_size == 0:
        raise ExcelValidationError(["El archivo está vacío (0 bytes)."])


def validate_header(header):
    errors = []
    columns = [c for c in header if not _is_blank(c)]
    if len(columns) != len(EXPECTED_COLUMNS):
        errors.append(
            f"Se esperaban {len(EXPECTED_COLUMNS)} columnas y se encontraron {len(columns)}."
        )
    normalized = tuple(str(c).strip().upper() for c in columns)
    if normalized != EXPECTED_COLUMNS:
        errors.append(
            f"Encabezados no válidos {list(normalized)}. Se esperaba {list(EXPECTED_COLUMNS)}."
        )
    return errors


def validate_row(row_number, row):
    errors = []
    doc_type, doc_number, salary = (list(row) + [None] * 3)[:3]

    if _is_blank(doc_type):
        errors.append(f"Fila {row_number}: TIPODOC está vacío.")
    elif str(doc_type).strip().upper() not in VALID_DOC_TYPES:
        errors.append(f"Fila {row_number}: TIPODOC '{doc_type}' no es válido.")

    if _is_blank(doc_number):
        errors.append(f"Fila {row_number}: NRODOCUMENTO está vacío.")
    elif not str(doc_number).strip().isalnum():
        errors.append(f"Fila {row_number}: NRODOCUMENTO '{doc_number}' no es válido.")

    if _is_blank(salary):
        errors.append(f"Fila {row_number}: SUELDO está vacío.")
    elif not _is_number(salary):
        errors.append(f"Fila {row_number}: SUELDO '{salary}' no es numérico.")
    elif float(salary) < 0:
        errors.append(f"Fila {row_number}: SUELDO no puede ser negativo.")

    return errors


def validate_excel(path):
    """Checks the file structure and returns its data rows as (row_number, row).

    Raises ExcelValidationError only for problems that prevent reading the file
    (missing file, wrong extension, empty sheet, bad header). Content problems
    such as invalid rows, repeated documents or notified types are handled by
    the reader, which skips them and reports warnings.
    """
    validate_file_path(path)

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, OSError, KeyError, ValueError) as exc:
        raise ExcelValidationError([f"No se pudo abrir el archivo: {exc}"]) from exc

    try:
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()

    # Keep the real Excel row number so messages match what the user sees
    rows = [(n, r) for n, r in enumerate(rows, start=1) if not all(_is_blank(v) for v in r)]
    if not rows:
        raise ExcelValidationError(["La hoja no tiene datos."])

    (_, header), data = rows[0], rows[1:]
    errors = validate_header(header)
    if errors:
        raise ExcelValidationError(errors)
    if not data:
        raise ExcelValidationError(["El archivo solo tiene encabezados, sin empleados."])

    return data
